from __future__ import annotations

import os
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

DEFAULT_WORKBOOK = Path('/Users/rasmusvikstrom/Desktop/ÖDT Tracker 2025 edition_2.xlsm')


def _safe_float(value: Any) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        cleaned = value.strip().replace(',', '.')
        try:
            return float(cleaned)
        except ValueError:
            return 0.0
    return float(value)


def _safe_iso(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat(timespec='seconds')
    return str(value)


def _find_workbook(workbook_path: str | None = None) -> Path | None:
    candidates: list[Path] = []
    if workbook_path:
        candidates.append(Path(workbook_path))
    env_path = os.getenv('ODT_WORKBOOK_PATH')
    if env_path:
        candidates.append(Path(env_path))
    candidates.append(DEFAULT_WORKBOOK)

    for candidate in candidates:
        if candidate.exists() and candidate.is_file():
            return candidate
    return None


def _parse_summary_rows(ws) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for row_idx in range(14, 27):
        name = ws.cell(row=row_idx, column=2).value
        if name is None or str(name).strip() == '':
            continue
        rows.append({
            'name': str(name),
            'total_volume': _safe_float(ws.cell(row=row_idx, column=3).value),
            'alcohol_amount': _safe_float(ws.cell(row=row_idx, column=4).value),
            'alcohol_per_kg': _safe_float(ws.cell(row=row_idx, column=5).value),
        })
    return rows


def _parse_metadata(ws) -> tuple[list[str], dict[str, int], dict[str, Any], dict[str, float]]:
    names: list[str] = []
    for col_idx in range(2, 30):
        value = ws.cell(row=2, column=col_idx).value
        if value is None or str(value).strip() == '':
            break
        names.append(str(value))

    team_map: dict[str, int] = {}
    for col_idx, name in zip(range(2, 2 + len(names)), names):
        value = ws.cell(row=3, column=col_idx).value
        team_map[name] = int(value) if value is not None else 0

    arrival_map: dict[str, Any] = {}
    for col_idx, name in zip(range(2, 2 + len(names)), names):
        arrival_map[name] = ws.cell(row=4, column=col_idx).value

    mass_map: dict[str, float] = {}
    for col_idx, name in zip(range(2, 2 + len(names)), names):
        mass_map[name] = _safe_float(ws.cell(row=5, column=col_idx).value)

    return names, team_map, arrival_map, mass_map


def _parse_team_standings(participants: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for participant in participants:
        grouped[str(participant['team'])].append(participant)

    rows: list[dict[str, Any]] = []
    for team, members in grouped.items():
        rows.append({
            'team': int(team),
            'members': [m['name'] for m in members],
            'total_volume': round(sum(m['total_volume'] for m in members), 2),
            'total_alcohol': round(sum(m['alcohol_amount'] for m in members), 2),
        })

    return sorted(rows, key=lambda row: row['total_alcohol'], reverse=True)


def _parse_timeline(wb) -> list[dict[str, Any]]:
    ws = wb['Analys']
    names_by_col: dict[int, str] = {}
    for col_idx in range(5, 30):
        value = ws.cell(row=31, column=col_idx).value
        if value is None or str(value).strip() == '':
            continue
        names_by_col[col_idx] = str(value)

    if not names_by_col:
        return []

    timeline: list[dict[str, Any]] = []
    for row_idx in range(33, ws.max_row + 1):
        timestamp_value = ws.cell(row=row_idx, column=4).value
        if timestamp_value is None:
            continue

        point: dict[str, Any] = {'timestamp': _safe_iso(timestamp_value)}
        has_data = False
        for col_idx, name in names_by_col.items():
            value = ws.cell(row=row_idx, column=col_idx).value
            numeric = _safe_float(value)
            point[name] = numeric
            if numeric != 0:
                has_data = True

        if has_data:
            timeline.append(point)

    return timeline


def load_tracker_data(workbook_path: str | None = None) -> dict[str, Any]:
    workbook_file = _find_workbook(workbook_path)
    if workbook_file is None:
        return {
            'status': 'missing_workbook',
            'message': 'No workbook found. Set ODT_WORKBOOK_PATH or place the Excel file on the default desktop path.',
            'participants': [],
            'team_standings': [],
            'timeline': [],
            'summary': {
                'title': 'ÖDT Tracker 2025',
                'event_start': None,
                'event_end': None,
                'leader': None,
                'team_leader': None,
            },
        }

    wb = load_workbook(workbook_file, data_only=False, read_only=True)
    input_ws = wb['Input']

    names, team_map, arrival_map, mass_map = _parse_metadata(input_ws)
    ranked_rows = _parse_summary_rows(input_ws)

    participants: list[dict[str, Any]] = []
    for rank_index, row in enumerate(ranked_rows, start=1):
        name = row['name']
        team_number = team_map.get(name, 0)
        participants.append({
            'rank': rank_index,
            'name': name,
            'team': team_number,
            'arrival': _safe_iso(arrival_map.get(name)),
            'mass': mass_map.get(name, 0.0),
            'total_volume': round(row['total_volume'], 2),
            'alcohol_amount': round(row['alcohol_amount'], 2),
            'alcohol_per_kg': round(row['alcohol_per_kg'], 2),
        })

    team_standings = _parse_team_standings(participants)
    timeline = _parse_timeline(wb)

    event_start = min((p['arrival'] for p in participants if p['arrival']), default=None)
    event_end = max((point['timestamp'] for point in timeline), default=None)
    leader = participants[0]['name'] if participants else None
    team_leader = team_standings[0]['team'] if team_standings else None

    summary = {
        'title': 'ÖDT Tracker 2025',
        'event_start': event_start,
        'event_end': event_end,
        'leader': leader,
        'team_leader': team_leader,
        'participant_count': len(participants),
        'team_count': len(team_standings),
        'top_total_volume': round(participants[0]['total_volume'], 2) if participants else 0.0,
        'top_alcohol_amount': round(participants[0]['alcohol_amount'], 2) if participants else 0.0,
    }

    return {
        'status': 'ok',
        'message': f'Loaded workbook: {workbook_file.name}',
        'participants': participants,
        'team_standings': team_standings,
        'timeline': timeline,
        'summary': summary,
    }
